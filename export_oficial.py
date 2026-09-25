import openpyxl
import sqlite3
import os
import io

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_FILE = os.path.join(BASE_DIR, 'template_alojamento.xlsx')
DB_FILE = os.path.join(BASE_DIR, 'alojamento.db')

def gerar_planilha_oficial():
    wb = openpyxl.load_workbook(TEMPLATE_FILE)
    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    # Mapear quartos e camas do banco de dados
    cursor.execute("""
    SELECT 
        b.nome as bloco_nome,
        q.numero as quarto_num,
        v.numero_cama,
        v.status,
        a.matricula,
        a.nome_completo as alojado_nome,
        a.funcao,
        e.nome as empresa_nome
    FROM vagas v
    JOIN quartos q ON v.quarto_id = q.id
    JOIN blocos b ON q.bloco_id = b.id
    LEFT JOIN alojados a ON a.vaga_id = v.id AND a.status = 'ativo'
    LEFT JOIN empresas e ON a.empresa_id = e.id
    ORDER BY b.nome, CAST(q.numero AS INTEGER), q.numero, v.numero_cama
    """)
    rows = cursor.fetchall()
    
    # Organizar por (bloco_nome, quarto_num) -> list of bed info
    dados_banco = {}
    for r in rows:
        b_name = r['bloco_nome']
        q_num = str(r['quarto_num'])
        key = (b_name, q_num)
        if key not in dados_banco:
            dados_banco[key] = []
        dados_banco[key].append({
            'status': r['status'],
            'matricula': r['matricula'] or '',
            'nome': r['alojado_nome'] or '',
            'funcao': r['funcao'] or '',
            'empresa': r['empresa_nome'] or ''
        })

    sheet_conteiner_name = wb.sheetnames[3] # Geralmente 'BLOCOS CONTÊINER'

    sections = [
        ('BLOCOS ALOJAMENTO', 10, 74, 'Bloco 1', 'Bloco 1', 
         {'q': 1, 'reg': 2, 'nome': 3, 'funcao': 8, 'emp': 11}, 
         {'q': 14, 'reg': 15, 'nome': 16, 'funcao': 21, 'emp': 24}),
        ('BLOCOS ALOJAMENTO', 84, 148, 'Bloco 2', 'Bloco 2', 
         {'q': 1, 'reg': 2, 'nome': 3, 'funcao': 8, 'emp': 11}, 
         {'q': 14, 'reg': 15, 'nome': 16, 'funcao': 21, 'emp': 24}),
        ('BLOCOS ALOJAMENTO', 159, 223, 'Bloco 3', 'Bloco 3', 
         {'q': 1, 'reg': 2, 'nome': 3, 'funcao': 8, 'emp': 11}, 
         {'q': 14, 'reg': 15, 'nome': 16, 'funcao': 21, 'emp': 24}),
        ('BLOCOS ALOJAMENTO', 242, 306, 'Bloco 4', 'Bloco 4', 
         {'q': 1, 'reg': 2, 'nome': 3, 'funcao': 8, 'emp': 11}, 
         {'q': 14, 'reg': 15, 'nome': 16, 'funcao': 21, 'emp': 24}),
        ('BLOCOS ALOJAMENTO', 317, 381, 'Bloco 5', 'Bloco 5', 
         {'q': 1, 'reg': 2, 'nome': 3, 'funcao': 8, 'emp': 11}, 
         {'q': 14, 'reg': 15, 'nome': 16, 'funcao': 21, 'emp': 24}),
        ('BLOCOS ADM', 10, 42, 'Bloco 1 ADM', 'Bloco 1 ADM', 
         {'q': 1, 'reg': 2, 'nome': 3, 'funcao': 8, 'emp': 11}, 
         {'q': 14, 'reg': 15, 'nome': 16, 'funcao': 21, 'emp': 24}),
        ('BLOCOS ADM', 57, 77, 'Bloco 2 ADM', 'Bloco 2 ADM', 
         {'q': 1, 'reg': 2, 'nome': 3, 'funcao': 8, 'emp': 11}, 
         {'q': 14, 'reg': 15, 'nome': 16, 'funcao': 21, 'emp': 24}),
        ('BLOCOS ADM', 90, 134, 'Bloco 3 ADM', 'Bloco 3 ADM', 
         {'q': 1, 'reg': 2, 'nome': 3, 'funcao': 8, 'emp': 11}, 
         {'q': 14, 'reg': 15, 'nome': 16, 'funcao': 21, 'emp': 24}),
        (sheet_conteiner_name, 9, 41, 'Contêiner 1', 'Contêiner 2', 
         {'q': 1, 'reg': 2, 'nome': 3, 'funcao': 8, 'emp': 11}, 
         {'q': 16, 'reg': 17, 'nome': 18, 'funcao': 23, 'emp': 26}),
    ]

    for sname, r_start, r_end, b_left, b_right, l_cols, r_cols in sections:
        ws = wb[sname]
        
        cur_q_l = None
        bed_idx_l = 0
        cur_q_r = None
        bed_idx_r = 0
        
        for r_idx in range(r_start, r_end):
            # openpyxl uses 1-based indexing for row
            row_excel = r_idx + 1
            
            # --- Lado Esquerdo ---
            ql_val = ws.cell(row=row_excel, column=l_cols['q'] + 1).value
            if ql_val is not None:
                ql_str = str(ql_val).strip()
                if ql_str and ql_str.replace('.', '').isdigit():
                    new_ql = str(int(float(ql_str)))
                    if new_ql != cur_q_l:
                        cur_q_l = new_ql
                        bed_idx_l = 0
            
            if cur_q_l:
                key_l = (b_left, cur_q_l)
                beds = dados_banco.get(key_l, [])
                if bed_idx_l < len(beds):
                    bed = beds[bed_idx_l]
                    if bed['status'] == 'ocupada' and bed['nome']:
                        ws.cell(row=row_excel, column=l_cols['reg'] + 1, value=bed['matricula'])
                        ws.cell(row=row_excel, column=l_cols['nome'] + 1, value=bed['nome'])
                        ws.cell(row=row_excel, column=l_cols['funcao'] + 1, value=bed['funcao'])
                        ws.cell(row=row_excel, column=l_cols['emp'] + 1, value=bed['empresa'])
                    else:
                        ws.cell(row=row_excel, column=l_cols['reg'] + 1, value="")
                        ws.cell(row=row_excel, column=l_cols['nome'] + 1, value="VAGA")
                        ws.cell(row=row_excel, column=l_cols['funcao'] + 1, value="")
                        ws.cell(row=row_excel, column=l_cols['emp'] + 1, value="")
                bed_idx_l += 1

            # --- Lado Direito ---
            if r_cols:
                qr_val = ws.cell(row=row_excel, column=r_cols['q'] + 1).value
                if qr_val is not None:
                    qr_str = str(qr_val).strip()
                    if qr_str and qr_str.replace('.', '').isdigit():
                        new_qr = str(int(float(qr_str)))
                        if new_qr != cur_q_r:
                            cur_q_r = new_qr
                            bed_idx_r = 0
                
                if cur_q_r:
                    key_r = (b_right, cur_q_r)
                    beds = dados_banco.get(key_r, [])
                    if bed_idx_r < len(beds):
                        bed = beds[bed_idx_r]
                        if bed['status'] == 'ocupada' and bed['nome']:
                            ws.cell(row=row_excel, column=r_cols['reg'] + 1, value=bed['matricula'])
                            ws.cell(row=row_excel, column=r_cols['nome'] + 1, value=bed['nome'])
                            ws.cell(row=row_excel, column=r_cols['funcao'] + 1, value=bed['funcao'])
                            ws.cell(row=row_excel, column=r_cols['emp'] + 1, value=bed['empresa'])
                        else:
                            ws.cell(row=row_excel, column=r_cols['reg'] + 1, value="")
                            ws.cell(row=row_excel, column=r_cols['nome'] + 1, value="VAGA")
                            ws.cell(row=row_excel, column=r_cols['funcao'] + 1, value="")
                            ws.cell(row=row_excel, column=r_cols['emp'] + 1, value="")
                    bed_idx_r += 1

    # Atualizar aba RESUMO (2) com totais consolidados
    try:
        ws_resumo = wb['RESUMO (2)']
        cursor.execute("SELECT COUNT(*) FROM vagas")
        total_vagas = cursor.fetchone()[0]
        cursor.execute("SELECT COUNT(*) FROM vagas WHERE status = 'ocupada'")
        total_ocupadas = cursor.fetchone()[0]
        total_livres = total_vagas - total_ocupadas
        
        # Row 5 (Produção Alojamentos)
        cursor.execute("""
        SELECT COUNT(*), SUM(CASE WHEN v.status='ocupada' THEN 1 ELSE 0 END)
        FROM vagas v JOIN quartos q ON v.quarto_id = q.id JOIN blocos b ON q.bloco_id = b.id
        WHERE b.tipo = 'alojamento'
        """)
        r_aloj = cursor.fetchone()
        tot_aloj = r_aloj[0] or 0
        ocu_aloj = r_aloj[1] or 0
        liv_aloj = tot_aloj - ocu_aloj
        
        ws_resumo.cell(row=5, column=5, value=tot_aloj)
        ws_resumo.cell(row=5, column=6, value=ocu_aloj)
        ws_resumo.cell(row=5, column=7, value=liv_aloj)
        
        # Row 6 (ADM & Contêineres)
        cursor.execute("""
        SELECT COUNT(*), SUM(CASE WHEN v.status='ocupada' THEN 1 ELSE 0 END)
        FROM vagas v JOIN quartos q ON v.quarto_id = q.id JOIN blocos b ON q.bloco_id = b.id
        WHERE b.tipo != 'alojamento'
        """)
        r_outros = cursor.fetchone()
        tot_outros = r_outros[0] or 0
        ocu_outros = r_outros[1] or 0
        liv_outros = tot_outros - ocu_outros
        
        ws_resumo.cell(row=6, column=5, value=tot_outros)
        ws_resumo.cell(row=6, column=6, value=ocu_outros)
        ws_resumo.cell(row=6, column=7, value=liv_outros)
    except Exception as e:
        print("Aviso ao atualizar resumo:", e)

    conn.close()
    
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

if __name__ == '__main__':
    buffer = gerar_planilha_oficial()
    print("Sucesso! Buffer gerado com tamanho:", len(buffer.getvalue()), "bytes")
