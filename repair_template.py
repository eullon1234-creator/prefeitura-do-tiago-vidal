import zipfile
import io
import openpyxl

orig_path = 'ALOJAMENTO TABOCA 2 (1) (4).xlsx'
template_path = 'template_alojamento.xlsx'

with open(orig_path, 'rb') as f:
    orig_bytes = f.read()

in_zip = zipfile.ZipFile(io.BytesIO(orig_bytes), 'r')
out_buffer = io.BytesIO()
out_zip = zipfile.ZipFile(out_buffer, 'w', zipfile.ZIP_DEFLATED)

for item in in_zip.infolist():
    data = in_zip.read(item.filename)
    if item.filename == 'xl/drawings/_rels/drawing3.xml.rels':
        data = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image1.png"/></Relationships>'
    out_zip.writestr(item, data)

out_zip.close()

with open(template_path, 'wb') as f:
    f.write(out_buffer.getvalue())

print(f"Template successfully written to {template_path}")

wb = openpyxl.load_workbook(template_path, data_only=False)
print("SUCCESS! Sheets loaded:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print(f"Sheet '{name}': rows={ws.max_row}, cols={ws.max_column}")
